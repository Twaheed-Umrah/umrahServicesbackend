import csv
import io
import logging
from django.db import transaction
from django.db.models import Q, Count
from rest_framework import generics, permissions, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser

from .models import Lead, LeadNote, Task
from .serializers import LeadSerializer, LeadUpdateStatusSerializer, TaskSerializer
from apps.users.models import User

logger = logging.getLogger(__name__)

class LeadListCreateView(generics.ListCreateAPIView):
    """
    List all leads filtered by role, or create a single lead.
    - Agency Head / SuperAdmin: views all agency leads.
    - Telecaller: views ONLY leads assigned to them.
    """
    serializer_class = LeadSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        queryset = Lead.objects.all().select_related('assigned_caller', 'user').prefetch_related('history_notes')

        if user.role == 'superadmin':
            tenant_id = self.request.query_params.get('tenant_id', None)
            if tenant_id and tenant_id != 'all':
                queryset = queryset.filter(Q(user_id=tenant_id) | Q(assigned_caller__created_by_id=tenant_id))
        elif user.role == 'sells_caller':
            # Callers only see their assigned leads
            queryset = queryset.filter(assigned_caller=user)
        else:
            # Agency Admin / Franchises Admin sees all their agency leads
            queryset = queryset.filter(Q(user=user) | Q(assigned_caller__created_by=user))

        # Query Filters
        lead_status = self.request.query_params.get('status', None)
        if lead_status and lead_status != 'all':
            queryset = queryset.filter(status=lead_status)

        caller_id = self.request.query_params.get('caller_id', None)
        if caller_id:
            queryset = queryset.filter(assigned_caller_id=caller_id)

        in_pool = self.request.query_params.get('in_pool', None)
        if in_pool is not None:
            if in_pool.lower() in ['true', '1']:
                queryset = queryset.filter(is_in_agency_pool=True)
            elif in_pool.lower() in ['false', '0']:
                queryset = queryset.filter(is_in_agency_pool=False)

        search = self.request.query_params.get('search', None)
        if search:
            queryset = queryset.filter(
                Q(name__icontains=search) |
                Q(mobile_number__icontains=search) |
                Q(email__icontains=search) |
                Q(city__icontains=search) |
                Q(campaign_name__icontains=search)
            )

        return queryset

    def perform_create(self, serializer):
        user = self.request.user
        agency_head = user if user.role != 'sells_caller' else user.created_by
        serializer.save(user=agency_head or user)


class LeadDetailView(generics.RetrieveUpdateDestroyAPIView):
    """
    Retrieve, update or delete a lead.
    """
    serializer_class = LeadSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.role == 'superadmin':
            return Lead.objects.all()
        elif user.role == 'sells_caller':
            return Lead.objects.filter(assigned_caller=user)
        else:
            return Lead.objects.filter(Q(user=user) | Q(assigned_caller__created_by=user))


class LeadAddNoteView(generics.UpdateAPIView):
    """
    Update lead pipeline status and append a call timeline note.
    """
    serializer_class = LeadUpdateStatusSerializer
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, *args, **kwargs):
        """Allow POST as an alias for partial update."""
        return self.partial_update(request, *args, **kwargs)

    def get_queryset(self):
        user = self.request.user
        if user.role == 'superadmin':
            return Lead.objects.all()
        elif user.role == 'sells_caller':
            return Lead.objects.filter(assigned_caller=user)
        else:
            return Lead.objects.filter(Q(user=user) | Q(assigned_caller__created_by=user))


class ExcelLeadUploadView(APIView):
    """
    Upload leads exported from Meta Ads Manager (Facebook/Instagram), Google Ads, or CSV/Excel.
    Automatically assigns leads in a cyclic Round-Robin sequence to all callers who are PRESENT.
    If no callers are present, safely stores leads in the Agency Head's Master Pool.
    """
    permission_classes = [permissions.IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        uploaded_file = request.FILES.get('file')
        if not uploaded_file:
            return Response({'error': 'No file was uploaded.'}, status=status.HTTP_400_BAD_REQUEST)

        filename = uploaded_file.name.lower()
        campaign_name = request.data.get('campaign_name', 'Meta Lead Ads')
        source = request.data.get('source', 'Meta Ads')

        extracted_rows = []

        try:
            if filename.endswith('.csv'):
                decoded_file = uploaded_file.read().decode('utf-8', errors='ignore')
                reader = csv.DictReader(io.StringIO(decoded_file))
                for row in reader:
                    extracted_rows.append(row)
            elif filename.endswith(('.xlsx', '.xls')):
                import openpyxl
                workbook = openpyxl.load_workbook(uploaded_file, data_only=True)
                sheet = workbook.active
                headers = [str(cell.value or '').strip().lower() for cell in next(sheet.iter_rows(max_row=1))]
                
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    row_dict = {}
                    for h_idx, header_name in enumerate(headers):
                        if h_idx < len(row) and header_name:
                            row_dict[header_name] = str(row[h_idx] or '').strip()
                    if any(row_dict.values()):
                        extracted_rows.append(row_dict)
            else:
                return Response({'error': 'Unsupported file format. Please upload a .csv or .xlsx file.'}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            logger.error(f"Error parsing uploaded lead file: {e}")
            return Response({'error': f'Failed to parse file: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)

        if not extracted_rows:
            return Response({'error': 'The uploaded spreadsheet contains no data rows.'}, status=status.HTTP_400_BAD_REQUEST)

        # Determine Agency Head
        agency_head = request.user if request.user.role != 'sells_caller' else request.user.created_by
        if not agency_head:
            agency_head = request.user

        # Fetch all active Callers for this agency who are PRESENT
        present_callers = list(User.objects.filter(
            created_by=agency_head,
            role='sells_caller',
            is_active=True,
            is_present=True  # Strictly exclude absent callers!
        ).order_by('id'))

        num_callers = len(present_callers)
        created_leads = []

        with transaction.atomic():
            for idx, raw in enumerate(extracted_rows):
                # Flexible header matching
                name = (
                    raw.get('full_name') or raw.get('name') or raw.get('full name') or
                    raw.get('lead_name') or raw.get('first_name') or raw.get('contact_name') or 'Pilgrim Lead'
                )
                mobile = (
                    raw.get('mobile_number') or raw.get('phone_number') or raw.get('phone') or
                    raw.get('mobile') or raw.get('contact') or raw.get('phone number') or ''
                )
                mobile = str(mobile).replace('.0', '').strip()
                if not mobile or len(mobile) < 7:
                    continue  # Invalid contact number

                email = raw.get('email') or raw.get('email_address') or raw.get('email address') or ''
                city = raw.get('city') or raw.get('location') or ''
                row_campaign = raw.get('campaign_name') or raw.get('campaign') or campaign_name

                # Round-Robin assignment
                if num_callers > 0:
                    assigned_caller = present_callers[idx % num_callers]
                    is_in_pool = False
                    assignment_note = f"Auto-assigned to {assigned_caller.get_full_name() or assigned_caller.username} via Round-Robin."
                else:
                    assigned_caller = None
                    is_in_pool = True
                    assignment_note = "Stored in Agency Master Pool (No callers currently marked as Present)."

                lead = Lead.objects.create(
                    user=agency_head,
                    assigned_caller=assigned_caller,
                    is_in_agency_pool=is_in_pool,
                    name=name,
                    mobile_number=mobile,
                    email=email or None,
                    city=city,
                    source=source,
                    campaign_name=row_campaign,
                    status='NEW',
                    notes=assignment_note
                )

                LeadNote.objects.create(
                    lead=lead,
                    note=assignment_note,
                    author=request.user
                )
                created_leads.append(lead)

        return Response({
            'success': True,
            'message': f"Successfully processed {len(created_leads)} leads.",
            'total_processed': len(created_leads),
            'present_callers_count': num_callers,
            'assigned_callers': [c.username for c in present_callers],
            'sent_to_master_pool': num_callers == 0
        }, status=status.HTTP_201_CREATED)


class LeadAssignSingleView(APIView):
    """
    Manually assign or reassign a single lead to a designated caller.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, lead_id):
        caller_id = request.data.get('caller_id')
        if not caller_id:
            return Response({'error': 'caller_id is required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            lead = Lead.objects.get(id=lead_id)
            caller = User.objects.get(id=caller_id, role='sells_caller', is_active=True)
        except Lead.DoesNotExist:
            return Response({'error': 'Lead not found.'}, status=status.HTTP_404_NOT_FOUND)
        except User.DoesNotExist:
            return Response({'error': 'Target caller not found or inactive.'}, status=status.HTTP_404_NOT_FOUND)

        old_caller = lead.assigned_caller.username if lead.assigned_caller else "Master Pool"
        lead.assigned_caller = caller
        lead.is_in_agency_pool = False
        lead.save()

        note_txt = f"Reassigned from {old_caller} to {caller.get_full_name() or caller.username} by {request.user.username}."
        LeadNote.objects.create(lead=lead, note=note_txt, author=request.user)

        return Response({
            'success': True,
            'message': f"Lead #{lead.id} successfully assigned to {caller.username}.",
            'lead': LeadSerializer(lead).data
        })


class LeadRecallToPoolView(APIView):
    """
    Pulls back leads from a selected caller into the Agency Head's master pool.
    Useful when a caller is absent, on leave, or non-performing.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        caller_id = request.data.get('caller_id')
        lead_ids = request.data.get('lead_ids', [])

        user = request.user
        agency_head = user if user.role != 'sells_caller' else user.created_by

        leads_to_recall = Lead.objects.filter(user=agency_head)

        if caller_id:
            try:
                caller = User.objects.get(id=caller_id)
            except User.DoesNotExist:
                return Response({'error': 'Caller not found.'}, status=status.HTTP_404_NOT_FOUND)

            # Pull all active/unconverted leads belonging to this caller
            leads_to_recall = leads_to_recall.filter(
                assigned_caller=caller,
                status__in=['NEW', 'CONTACTED', 'FOLLOW_UP', 'BUSY', 'RNR', 'CALLBACK']
            )
            caller_name = caller.get_full_name() or caller.username
        elif lead_ids:
            leads_to_recall = leads_to_recall.filter(id__in=lead_ids)
            caller_name = "selected callers"
        else:
            return Response({'error': 'Provide either caller_id or lead_ids list.'}, status=status.HTTP_400_BAD_REQUEST)

        count = leads_to_recall.count()
        if count == 0:
            return Response({'success': True, 'message': 'No eligible active leads to recall.', 'recalled_count': 0})

        with transaction.atomic():
            for lead in leads_to_recall:
                prev_caller = lead.assigned_caller.username if lead.assigned_caller else "Unknown"
                lead.assigned_caller = None
                lead.is_in_agency_pool = True
                lead.save(update_fields=['assigned_caller', 'is_in_agency_pool', 'updated_at'])

                LeadNote.objects.create(
                    lead=lead,
                    note=f"Recalled from {prev_caller} to Agency Master Pool by {request.user.username}.",
                    author=request.user
                )

        return Response({
            'success': True,
            'message': f"Successfully recalled {count} leads from {caller_name} to the Master Pool.",
            'recalled_count': count
        })


class LeadReassignFromPoolView(APIView):
    """
    Reassigns leads from the master pool to a target caller, or distributes them via Round-Robin.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        lead_ids = request.data.get('lead_ids', [])
        target_caller_id = request.data.get('target_caller_id')
        use_round_robin = request.data.get('round_robin', False)

        user = request.user
        agency_head = user if user.role != 'sells_caller' else user.created_by

        pool_leads = Lead.objects.filter(user=agency_head, is_in_agency_pool=True)
        if lead_ids:
            pool_leads = pool_leads.filter(id__in=lead_ids)

        if not pool_leads.exists():
            return Response({'error': 'No eligible pool leads found.'}, status=status.HTTP_400_BAD_REQUEST)

        if target_caller_id:
            try:
                target_caller = User.objects.get(id=target_caller_id, role='sells_caller', is_active=True)
            except User.DoesNotExist:
                return Response({'error': 'Target caller not found or inactive.'}, status=status.HTTP_404_NOT_FOUND)

            count = pool_leads.count()
            with transaction.atomic():
                for lead in pool_leads:
                    lead.assigned_caller = target_caller
                    lead.is_in_agency_pool = False
                    lead.save(update_fields=['assigned_caller', 'is_in_agency_pool', 'updated_at'])

                    LeadNote.objects.create(
                        lead=lead,
                        note=f"Assigned from Master Pool to {target_caller.username} by {request.user.username}.",
                        author=request.user
                    )

            return Response({
                'success': True,
                'message': f"Assigned {count} leads from Master Pool to {target_caller.username}.",
                'assigned_count': count
            })

        elif use_round_robin:
            present_callers = list(User.objects.filter(
                created_by=agency_head,
                role='sells_caller',
                is_active=True,
                is_present=True
            ).order_by('id'))

            if not present_callers:
                return Response({'error': 'Cannot distribute: No callers are currently marked as Present.'}, status=status.HTTP_400_BAD_REQUEST)

            count = 0
            with transaction.atomic():
                for idx, lead in enumerate(pool_leads):
                    caller = present_callers[idx % len(present_callers)]
                    lead.assigned_caller = caller
                    lead.is_in_agency_pool = False
                    lead.save(update_fields=['assigned_caller', 'is_in_agency_pool', 'updated_at'])

                    LeadNote.objects.create(
                        lead=lead,
                        note=f"Redistributed from Master Pool to {caller.username} via Round-Robin.",
                        author=request.user
                    )
                    count += 1

            return Response({
                'success': True,
                'message': f"Distributed {count} leads across {len(present_callers)} present callers via Round-Robin.",
                'distributed_count': count
            })
        else:
            return Response({'error': 'Specify target_caller_id or set round_robin=True.'}, status=status.HTTP_400_BAD_REQUEST)


class LeadAgencyPoolListView(generics.ListAPIView):
    """
    List all leads sitting in the Agency Head's master unassigned pool.
    """
    serializer_class = LeadSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        agency_head = user if user.role != 'sells_caller' else user.created_by
        return Lead.objects.filter(
            Q(user=agency_head) & (Q(is_in_agency_pool=True) | Q(assigned_caller__isnull=True))
        ).order_by('-created_at')


class TaskListCreateView(generics.ListCreateAPIView):
    """Module 18: Task & Team Management"""
    serializer_class = TaskSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.role == 'superadmin':
            return Task.objects.all().select_related('assigned_to', 'created_by', 'related_lead')
        elif user.role in ['franchisesadmin', 'franchiseadmin']:
            return Task.objects.filter(
                Q(created_by=user) | Q(assigned_to=user) | Q(assigned_to__created_by=user)
            ).select_related('assigned_to', 'created_by', 'related_lead')
        else:
            return Task.objects.filter(assigned_to=user).select_related('assigned_to', 'created_by', 'related_lead')


class TaskDetailView(generics.RetrieveUpdateDestroyAPIView):
    serializer_class = TaskSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.role == 'superadmin':
            return Task.objects.all()
        return Task.objects.filter(
            Q(created_by=user) | Q(assigned_to=user) | Q(assigned_to__created_by=user)
        )
